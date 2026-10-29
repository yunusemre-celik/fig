import sys
import os
import pypdf

def pdf_to_excel(pdf_path, out_path):
    import openpyxl
    reader = pypdf.PdfReader(pdf_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PDF Tablosu"
    
    current_row = 1
    for page_idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        ws.cell(row=current_row, column=1, value=f"[Sayfa {page_idx + 1}]")
        current_row += 1
        
        for line in text.splitlines():
            line_str = line.strip()
            if not line_str:
                continue
            # Try splitting by tabs first, then multiple spaces
            parts = [p.strip() for p in line.split('\t') if p.strip()]
            if len(parts) <= 1:
                parts = [p.strip() for p in line.split('   ') if p.strip()]
            if not parts:
                parts = [line_str]
                
            for col_idx, part in enumerate(parts, start=1):
                ws.cell(row=current_row, column=col_idx, value=part)
            current_row += 1
        current_row += 1
        
    wb.save(out_path)
    print(f"SUCCESS: {out_path}")

def pdf_to_docx(pdf_path, out_path):
    import docx
    reader = pypdf.PdfReader(pdf_path)
    doc = docx.Document()
    
    for page_idx, page in enumerate(reader.pages):
        if page_idx > 0:
            doc.add_page_break()
        doc.add_heading(f"Sayfa {page_idx + 1}", level=2)
        text = page.extract_text() or ""
        for line in text.splitlines():
            line_str = line.strip()
            if line_str:
                doc.add_paragraph(line_str)
                
    doc.save(out_path)
    print(f"SUCCESS: {out_path}")

def pdf_to_text(pdf_path, out_path):
    reader = pypdf.PdfReader(pdf_path)
    all_text = []
    for page_idx, page in enumerate(reader.pages):
        all_text.append(f"--- Sayfa {page_idx + 1} ---")
        text = page.extract_text() or ""
        all_text.append(text)
    
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(all_text))
    print(f"SUCCESS: {out_path}")

def pdf_to_images(pdf_path, out_dir):
    reader = pypdf.PdfReader(pdf_path)
    os.makedirs(out_dir, exist_ok=True)
    count = 0
    for page_idx, page in enumerate(reader.pages):
        for img_idx, img in enumerate(page.images):
            count += 1
            out_img = os.path.join(out_dir, f"sayfa_{page_idx+1}_gorsel_{img_idx+1}.png")
            with open(out_img, "wb") as f:
                f.write(img.data)
                
    if count == 0:
        # If no embedded images, write text summary
        summary_path = os.path.join(out_dir, "pdf_icerik.txt")
        pdf_to_text(pdf_path, summary_path)
    print(f"SUCCESS: {count} images extracted to {out_dir}")

def main():
    if len(sys.argv) < 4:
        print("Usage: python pdf_converter.py <to_excel|to_docx|to_text|to_images> <input_pdf> <output_path>")
        sys.exit(1)
        
    cmd = sys.argv[1].lower()
    in_pdf = sys.argv[2]
    out_target = sys.argv[3]
    
    if not os.path.exists(in_pdf):
        print(f"ERROR: Input file does not exist: {in_pdf}")
        sys.exit(1)
        
    try:
        if cmd == "to_excel":
            pdf_to_excel(in_pdf, out_target)
        elif cmd == "to_docx":
            pdf_to_docx(in_pdf, out_target)
        elif cmd == "to_text":
            pdf_to_text(in_pdf, out_target)
        elif cmd == "to_images":
            pdf_to_images(in_pdf, out_target)
        else:
            print(f"ERROR: Unknown command {cmd}")
            sys.exit(1)
    except Exception as e:
        print(f"ERROR: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()
